from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation


OUTPUT = Path(__file__).with_name("Template_Ujian.xlsx")
NAVY = "17365D"
WHITE = "FFFFFF"
YELLOW = "FFF2CC"

STUDENTS = ["Kelas", "No Absen", "Nama"]
EXAMS = [
    "Kode", "Nama Ujian", "Token", "Kelas", "Buka", "Tutup",
    "Durasi (menit)", "Tampilkan Nilai", "Tampilkan Pembahasan",
    "Maks Pelanggaran", "Jeda (detik)", "Acak Soal", "Acak Opsi", "Status",
]
QUESTIONS = [
    "No", "Tipe", "Pertanyaan", "Gambar", "A", "B", "C", "D", "E",
    "Kunci", "Bobot", "Pembahasan",
]
PARTICIPANTS = [
    "Kelas", "No Absen", "Nama", "Status", "Login Terakhir", "Mulai",
    "Selesai", "Pelanggaran", "Riwayat Pelanggaran", "Skor Otomatis",
    "Skor Manual", "Total Skor", "Skor Maks", "Nilai", "Sesi",
    "Jawaban Sementara", "Data",
]
DETAIL_ANSWERS = [
    "Waktu", "Kelas", "No Absen", "Nama", "No", "Tipe", "Pertanyaan",
    "Jawaban", "Kunci", "Benar", "Skor", "Bobot",
]
SAMPLE_QUESTIONS = [
    [1, "PG", r"Hitung nilai $\frac{1}{2}+\frac{1}{4}$.", "", r"$\frac{1}{2}$", r"$\frac{3}{4}$", "$1$", r"$\frac{1}{4}$", "", "B", 2, r"Samakan penyebut menjadi 4: $\frac{2}{4}+\frac{1}{4}=\frac{3}{4}$."],
    [2, "PG", r"Jika $2x+3=11$, nilai $x$ adalah …", "", "2", "3", "4", "5", "", "C", 2, "Kurangi 3 dari kedua sisi, lalu bagi 2."],
    [3, "PG", r"Sederhanakan $3a+2a-4$.", "", "$5a-4$", "$a-4$", "$5a+4$", "$6a$", "", "A", 2, "Gabungkan suku sejenis: $3a+2a=5a$."],
    [4, "ISIAN", r"Hitung $5^2-3^2$.", "", "", "", "", "", "", "16", 2, "Karena $25-9=16$."],
    [5, "ISIAN", r"Tuliskan $0{,}75$ sebagai pecahan.", "", "", "", "", "", "", "3/4|0,75", 2, r"$0{,}75=\frac{75}{100}=\frac{3}{4}$."],
    [6, "ESSAY", r"Sebuah persegi panjang panjangnya $8$ cm dan lebarnya $5$ cm. Hitung keliling dan luasnya.", "", "", "", "", "", "", "", 4, r"Keliling $=2(p+l)=26$ cm dan luas $=p\times l=40$ cm²."],
]
WIDE_ANSWERS = ["Waktu Selesai", "Kelas", "No Absen", "Nama"]
for sample_question in SAMPLE_QUESTIONS:
    question_number, question_type = sample_question[0], sample_question[1]
    WIDE_ANSWERS.append(f"Jawaban {question_number}")
    if question_type != "PG":
        WIDE_ANSWERS.append(f"Nilai Guru {question_number}")
WIDE_ANSWERS.extend(["Skor PG", "Nilai Akumulasi (/100)"])


def setup_sheet(ws, headers, widths=None):
    ws.append(headers)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{ws.cell(1, len(headers)).coordinate}"
    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = Font(color=WHITE, bold=True)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 30
    for index, header in enumerate(headers, start=1):
        ws.column_dimensions[ws.cell(1, index).column_letter].width = (
            widths[index - 1] if widths and index <= len(widths) else max(14, min(30, len(header) + 3))
        )


def add_list_validation(ws, formula, cell_range):
    validation = DataValidation(type="list", formula1=formula, allow_blank=True)
    validation.error = "Pilih nilai dari daftar."
    validation.errorTitle = "Pilihan tidak valid"
    validation.showErrorMessage = True
    ws.add_data_validation(validation)
    validation.add(cell_range)


def make_workbook():
    wb = Workbook()
    students = wb.active
    students.title = "Siswa"
    setup_sheet(students, STUDENTS, [14, 14, 28])
    for row in [["8A", 1, "Andi Pratama"], ["8A", 2, "Bunga Lestari"], ["8B", 1, "Citra Maharani"], ["8B", 2, "Dimas Saputra"]]:
        students.append(row)
    for row in range(2, 1001):
        students[f"B{row}"].number_format = "0"

    exams = wb.create_sheet("Daftar Ujian")
    setup_sheet(exams, EXAMS, [13, 25, 14, 16, 20, 20, 16, 20, 23, 19, 14, 13, 13, 17])
    exams.append(["KUIS1", "Kuis Matematika 1", "MTK8", "8A,8B", None, None, 20, "Langsung", "Ya", 3, 60, "Ya", "Ya", "Aktif"])
    for row in range(2, 1001):
        for col in ("E", "F"):
            exams[f"{col}{row}"].number_format = "yyyy-mm-dd hh:mm"
        exams[f"G{row}"].number_format = "0"
    for col in ("H",):
        add_list_validation(exams, '"Tidak,Langsung,Setelah ditutup"', f"{col}2:{col}1000")
    for col in ("I", "L", "M"):
        add_list_validation(exams, '"Ya,Tidak"', f"{col}2:{col}1000")
    add_list_validation(exams, '"Aktif,Nonaktif,Diarsipkan"', "N2:N1000")
    for row in exams.iter_rows(min_row=2, min_col=3, max_col=3):
        for cell in row:
            cell.number_format = "@"

    questions = wb.create_sheet("Soal_KUIS1")
    setup_sheet(questions, QUESTIONS, [9, 13, 55, 30, 22, 22, 22, 22, 22, 18, 12, 55])
    for row in SAMPLE_QUESTIONS:
        questions.append(row)
    add_list_validation(questions, '"PG,ISIAN,ESSAY"', "B2:B1000")
    for col in "DEFGHIJ":
        for row in range(2, max(1000, questions.max_row) + 1):
            questions[f"{col}{row}"].number_format = "@"
    questions.column_dimensions["D"].width = 35
    for row in questions.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    participants = wb.create_sheet("Peserta_KUIS1")
    setup_sheet(participants, PARTICIPANTS, [13, 12, 24, 16, 20, 20, 20, 14, 36, 17, 15, 14, 14, 13, 36, 36, 36])
    for row in range(2, 1001):
        participants[f"B{row}"].number_format = "0"
        for col in ("E", "F", "G"):
            participants[f"{col}{row}"].number_format = "yyyy-mm-dd hh:mm:ss"
        for col in ("J", "K", "L", "M", "N"):
            participants[f"{col}{row}"].number_format = "0.0"
    for col in ("O", "P", "Q"):
        participants.column_dimensions[col].hidden = True
    add_list_validation(participants, '"Login,Mengerjakan,Selesai"', "D2:D1000")

    answers = wb.create_sheet("Jawaban_KUIS1")
    setup_sheet(answers, WIDE_ANSWERS, [21, 13, 12, 24] + [28 if header.startswith("Jawaban ") else 16 for header in WIDE_ANSWERS[4:]])
    answer_question = {str(row[0]): row for row in SAMPLE_QUESTIONS}
    for col_index, header in enumerate(WIDE_ANSWERS, start=1):
        if header.startswith("Jawaban "):
            for row in range(2, 1001):
                answers.cell(row, col_index).number_format = "@"
        elif header in ("Skor PG", "Nilai Akumulasi (/100)"):
            for row in range(2, 1001):
                answers.cell(row, col_index).number_format = "0.0"
        elif header.startswith("Nilai Guru "):
            question_number = header.removeprefix("Nilai Guru ")
            question = answer_question[question_number]
            for row in range(2, 1001):
                answers.cell(row, col_index).fill = PatternFill("solid", fgColor=YELLOW)
                answers.cell(row, col_index).number_format = "0.0"
            validation = DataValidation(
                type="decimal", operator="between", formula1="0",
                formula2=str(question[10]), allow_blank=True,
            )
            validation.error = "Nilai harus berada di antara 0 dan bobot soal."
            validation.errorTitle = "Skor di luar batas"
            validation.showErrorMessage = True
            answers.add_data_validation(validation)
            validation.add(f"{answers.cell(2, col_index).column_letter}2:{answers.cell(1000, col_index).column_letter}1000")
    for row in range(2, 1001):
        answers[f"C{row}"].number_format = "0"
        answers[f"A{row}"].number_format = "yyyy-mm-dd hh:mm:ss"

    detail_answers = wb.create_sheet("Detail_KUIS1")
    setup_sheet(detail_answers, DETAIL_ANSWERS, [21, 13, 12, 24, 10, 13, 52, 35, 22, 12, 13, 12])
    for row in range(2, 1001):
        detail_answers.cell(row, 11).fill = PatternFill("solid", fgColor=YELLOW)
    for row in detail_answers.iter_rows(min_row=2, min_col=8, max_col=9, max_row=1000):
        for cell in row:
            cell.number_format = "@"
    add_list_validation(detail_answers, '"PG,ISIAN,ESSAY"', "F2:F1000")
    add_list_validation(detail_answers, '"Ya,Tidak"', "J2:J1000")
    wb._sheets.remove(answers)
    wb._sheets.insert(wb._sheets.index(detail_answers) + 1, answers)

    guide = wb.create_sheet("Petunjuk")
    setup_sheet(guide, ["Sheet", "Kolom", "Petunjuk"], [24, 28, 100])
    wide_answer_descriptions = [
        "Waktu pengumpulan siswa.", "Kelas peserta.", "Nomor absen peserta.", "Nama peserta.",
    ]
    for sample_question in SAMPLE_QUESTIONS:
        wide_answer_descriptions.append("Jawaban untuk soal sesuai nomor pada judul kolom.")
        if sample_question[1] != "PG":
            wide_answer_descriptions.append("Nilai ISIAN/ESSAY yang diisi guru; warna kuning, dari 0 sampai bobot soal.")
    wide_answer_descriptions.extend([
        "Jumlah skor PG yang dinilai otomatis.", "Nilai sementara dari akumulasi semua skor yang sudah tersedia, dalam skala 100.",
    ])
    descriptions = {
        "Siswa": ["Nama kelas siswa.", "Nomor absen; dibandingkan sebagai angka.", "Nama lengkap siswa."],
        "Daftar Ujian": [
            "Kode unik huruf/angka tanpa spasi.", "Nama ujian yang tampil ke siswa.", "Token masuk; tidak peka huruf besar/kecil.",
            "Kelas yang boleh ikut, dipisah koma, atau SEMUA.", "Waktu mulai ujian; kosong berarti tanpa batas awal.",
            "Waktu tutup; kosong berarti tanpa batas akhir.", "Durasi per siswa dalam menit.", "Aturan menampilkan nilai.",
            "Apakah pembahasan ditampilkan.", "Jumlah pelanggaran maksimum; 0 berarti tidak dibatasi.",
            "Jeda setelah pelanggaran, dalam detik.", "Acak urutan soal (Ya/Tidak).", "Acak urutan opsi PG (Ya/Tidak).",
            "Status ujian.",
        ],
        "Soal_KUIS1": [
            "Nomor soal.", "PG, ISIAN, atau ESSAY.", "Teks soal; mendukung rumus LaTeX.", "URL gambar atau Google Drive.",
            "Teks opsi A.", "Teks opsi B.", "Teks opsi C.", "Teks opsi D.", "Teks opsi E.",
            "Kunci: huruf untuk PG; opsional sebagai referensi guru untuk ISIAN; boleh kosong untuk ESSAY.",
            "Poin maksimal soal.", "Pembahasan yang dapat ditampilkan setelah ujian.",
        ],
        "Peserta_KUIS1": [
            "Kelas peserta.", "Nomor absen.", "Nama peserta.", "Status pengerjaan.", "Waktu login terakhir.",
            "Waktu mulai.", "Waktu selesai.", "Jumlah pelanggaran.", "Riwayat waktu dan jenis pelanggaran.",
            "Jumlah skor PG otomatis, dihitung aplikasi.", "Jumlah nilai ISIAN/ESSAY dari guru, dihitung aplikasi.", "Akumulasi total skor dari aplikasi.", "Jumlah semua bobot soal.",
            "Nilai dalam skala 0–100, dihitung aplikasi.", "ID sesi aktif (tersembunyi).", "Jawaban sementara JSON (tersembunyi).",
            "Urutan soal/opsi dan deadline JSON (tersembunyi).",
        ],
        "Jawaban_KUIS1": wide_answer_descriptions,
        "Detail_KUIS1": [
            "Waktu jawaban dicatat.", "Kelas peserta.", "Nomor absen peserta.", "Nama peserta.", "Nomor soal.",
            "Jenis soal.", "Teks soal.", "Jawaban siswa.", "Kunci jawaban.", "Hasil otomatis Ya/Tidak.",
            "Skor per soal; nilai ISIAN/ESSAY disalin dari kolom kuning di Jawaban_KUIS1.", "Bobot maksimal soal.",
        ],
    }
    for sheet_name, headers in [
        ("Siswa", STUDENTS), ("Daftar Ujian", EXAMS), ("Soal_KUIS1", QUESTIONS),
        ("Peserta_KUIS1", PARTICIPANTS), ("Jawaban_KUIS1", WIDE_ANSWERS),
        ("Detail_KUIS1", DETAIL_ANSWERS),
    ]:
        for header, description in zip(headers, descriptions[sheet_name]):
            guide.append([sheet_name, header, description])
    for row in guide.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    guide.auto_filter.ref = f"A1:C{guide.max_row}"
    wb.save(OUTPUT)


if __name__ == "__main__":
    make_workbook()
    print(f"Template berhasil dibuat: {OUTPUT}")
